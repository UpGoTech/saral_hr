import frappe
from frappe.utils import getdate, flt, add_months


MONTH_MAP = {
    'January':1,'February':2,'March':3,'April':4,
    'May':5,'June':6,'July':7,'August':8,
    'September':9,'October':10,'November':11,'December':12
}
MONTH_NUM_TO_NAME = {v:k for k,v in MONTH_MAP.items()}


def _start_date_for(year, month_num):
    return f"{year}-{month_num:02d}-01"


def _get_permitted_employees():
    """
    Returns a list of employee IDs the current user is permitted to see,
    based on Frappe User Permissions.

    - If the user has Employee-level User Permissions set → return only those employee IDs.
    - If NO Employee-level permissions are set → return None (meaning: show all, no restriction).

    This means Managers/Admins with no Employee permissions see everyone,
    while restricted users (like Sagar) only see their permitted employees.
    """
    user = frappe.session.user

    # System Manager / Administrator always sees everything
    if "System Manager" in frappe.get_roles(user):
        return None

    user_permissions = frappe.permissions.get_user_permissions(user)

    # Check if Employee-level restrictions exist
    if "Employee" in user_permissions and user_permissions["Employee"]:
        permitted = [perm.get("doc") for perm in user_permissions["Employee"] if perm.get("doc")]
        return permitted if permitted else None

    return None  # No restriction — show all


def _slips_for_period(company, start_date, permitted_employees=None):
    """
    Fetch salary slips for a given company + period.
    If permitted_employees is a list, only fetch slips for those employees.
    If None, fetch all (no restriction).
    """
    if permitted_employees is not None and len(permitted_employees) == 0:
        # User has permissions configured but 0 employees allowed — return empty
        return []

    if permitted_employees is not None:
        # Build IN clause safely using frappe.db.sql with tuple
        return frappe.db.sql("""
            SELECT
                ss.name,
                ss.employee,
                ss.employee_name,
                ss.docstatus,
                ss.net_salary,
                ss.total_earnings,
                ss.total_deductions,
                ss.total_employer_contribution,
                ss.payment_days,
                ss.category,
                ss.department,
                ss.designation,
                ss.branch,
                ss.division
            FROM `tabSalary Slip` ss
            WHERE ss.start_date = %(start_date)s
              AND ss.company    = %(company)s
              AND ss.docstatus  IN (0, 1, 2)
              AND ss.employee   IN %(permitted)s
            ORDER BY ss.employee_name ASC
        """, {
            "company": company,
            "start_date": start_date,
            "permitted": tuple(permitted_employees)
        }, as_dict=1)
    else:
        return frappe.db.sql("""
            SELECT
                ss.name,
                ss.employee,
                ss.employee_name,
                ss.docstatus,
                ss.net_salary,
                ss.total_earnings,
                ss.total_deductions,
                ss.total_employer_contribution,
                ss.payment_days,
                ss.category,
                ss.department,
                ss.designation,
                ss.branch,
                ss.division
            FROM `tabSalary Slip` ss
            WHERE ss.start_date = %(start_date)s
              AND ss.company    = %(company)s
              AND ss.docstatus  IN (0, 1, 2)
            ORDER BY ss.employee_name ASC
        """, {"company": company, "start_date": start_date}, as_dict=1)


@frappe.whitelist()
def get_salary_insight_data(company, year, month):
    month_num = MONTH_MAP.get(month)
    if not month_num:
        frappe.throw("Invalid month")

    year = int(year)
    start_date = _start_date_for(year, month_num)

    # ── Determine permitted employees for this user ───────────────────────────
    permitted_employees = _get_permitted_employees()

    # ── Active employees ──────────────────────────────────────────────────────
    # If user is restricted, count only their permitted employees
    if permitted_employees is not None:
        if len(permitted_employees) == 0:
            total_active = 0
        else:
            total_active = frappe.db.count(
                "Company Link",
                {
                    "company": company,
                    "is_active": 1,
                    "name": ["in", permitted_employees]
                }
            )
    else:
        total_active = frappe.db.count("Company Link", {"company": company, "is_active": 1})

    # ── Salary slips for this period ──────────────────────────────────────────
    salary_slips = _slips_for_period(company, start_date, permitted_employees)

    # ── Aggregated totals (submitted only) ────────────────────────────────────
    def aggregate(slips):
        te = td = tn = tec = 0
        for s in slips:
            if s.docstatus == 1:
                te  += flt(s.total_earnings)
                td  += flt(s.total_deductions)
                tn  += flt(s.net_salary)
                tec += flt(s.total_employer_contribution)
        return {"total_earnings": flt(te,2), "total_deductions": flt(td,2),
                "total_net": flt(tn,2), "total_employer_contribution": flt(tec,2)}

    totals = aggregate(salary_slips)

    # ── Previous month totals ─────────────────────────────────────────────────
    prev_date_obj  = getdate(add_months(start_date, -1))
    prev_start     = _start_date_for(prev_date_obj.year, prev_date_obj.month)
    prev_slips     = _slips_for_period(company, prev_start, permitted_employees)
    prev_totals    = aggregate(prev_slips)

    # ── Trend: last 6 months ──────────────────────────────────────────────────
    trend = []
    for i in range(5, -1, -1):
        t_date  = getdate(add_months(start_date, -i))
        t_start = _start_date_for(t_date.year, t_date.month)
        t_slips = _slips_for_period(company, t_start, permitted_employees)
        t_net   = sum(flt(s.net_salary) for s in t_slips if s.docstatus == 1)
        trend.append({
            "month": MONTH_NUM_TO_NAME[t_date.month],
            "year":  t_date.year,
            "net":   flt(t_net, 2),
        })

    # ── Category summary ──────────────────────────────────────────────────────
    if permitted_employees is not None:
        if len(permitted_employees) == 0:
            active_by_cat = []
        else:
            active_by_cat = frappe.db.sql("""
                SELECT category, COUNT(*) AS cnt
                FROM `tabCompany Link`
                WHERE company = %(company)s
                  AND is_active = 1
                  AND name IN %(permitted)s
                GROUP BY category
            """, {"company": company, "permitted": tuple(permitted_employees)}, as_dict=1)
    else:
        active_by_cat = frappe.db.sql("""
            SELECT category, COUNT(*) AS cnt
            FROM `tabCompany Link`
            WHERE company = %(company)s AND is_active = 1
            GROUP BY category
        """, {"company": company}, as_dict=1)

    cat_active_map = {r.category: r.cnt for r in active_by_cat}

    cat_map = {}
    for s in salary_slips:
        cat = s.category or "Uncategorised"
        if cat not in cat_map:
            cat_map[cat] = {"category": cat, "total": 0, "submitted": 0,
                            "draft": 0, "cancelled": 0, "active": 0, "net_amount": 0}
        cat_map[cat]["total"] += 1
        if s.docstatus == 1:
            cat_map[cat]["submitted"]  += 1
            cat_map[cat]["net_amount"] += flt(s.net_salary)
        elif s.docstatus == 0:
            cat_map[cat]["draft"] += 1
        elif s.docstatus == 2:
            cat_map[cat]["cancelled"] += 1

    for cat_name, cnt in cat_active_map.items():
        key = cat_name or "Uncategorised"
        if key not in cat_map:
            cat_map[key] = {"category": key, "total": 0, "submitted": 0,
                            "draft": 0, "cancelled": 0, "active": cnt, "net_amount": 0}
        else:
            cat_map[key]["active"] = cnt

    category_summary = sorted(cat_map.values(), key=lambda x: x["net_amount"], reverse=True)

    # ── Employees with NO salary slip ─────────────────────────────────────────
    slip_employee_set = {s.employee for s in salary_slips if s.docstatus != 2}

    if permitted_employees is not None:
        if len(permitted_employees) == 0:
            all_active_emps = []
        else:
            all_active_emps = frappe.db.sql("""
                SELECT cl.name, cl.full_name AS employee_name,
                       cl.category, cl.department, cl.designation
                FROM `tabCompany Link` cl
                WHERE cl.company = %(company)s
                  AND cl.is_active = 1
                  AND cl.name IN %(permitted)s
                ORDER BY cl.full_name ASC
            """, {"company": company, "permitted": tuple(permitted_employees)}, as_dict=1)
    else:
        all_active_emps = frappe.db.sql("""
            SELECT cl.name, cl.full_name AS employee_name,
                   cl.category, cl.department, cl.designation
            FROM `tabCompany Link` cl
            WHERE cl.company = %(company)s AND cl.is_active = 1
            ORDER BY cl.full_name ASC
        """, {"company": company}, as_dict=1)

    pending_employees = [e for e in all_active_emps if e.name not in slip_employee_set]

    return {
        "salary_slips":           salary_slips,
        "totals":                 totals,
        "prev_totals":            prev_totals,
        "category_summary":       category_summary,
        "total_active_employees": total_active,
        "pending_employees":      pending_employees,
        "trend":                  trend,
    }


@frappe.whitelist()
def get_all_slip_deductions(slip_names):
    import json
    if isinstance(slip_names, str):
        slip_names = json.loads(slip_names)

    # ── Guard: nothing to process ─────────────────────────────────────────────
    if not slip_names:
        return {}

    # ── Security: filter slip_names to only those the user can access ─────────
    permitted_employees = _get_permitted_employees()

    if permitted_employees is not None:
        if len(permitted_employees) == 0:
            return {}
        # tuple(slip_names) is safe here because we checked len > 0 above
        allowed_slips = frappe.db.sql("""
            SELECT name FROM `tabSalary Slip`
            WHERE name IN %(names)s
              AND employee IN %(permitted)s
        """, {
            "names":     tuple(slip_names),
            "permitted": tuple(permitted_employees)
        }, as_dict=1)
        slip_names = [r.name for r in allowed_slips]

    # ── Guard: nothing left after permission filter ───────────────────────────
    if not slip_names:
        return {}

    result = {}

    for name in slip_names:
        try:
            doc = frappe.get_doc("Salary Slip", name)

            if doc.docstatus != 1:
                continue

            # ── Attendance ────────────────────────────────────────────────
            total_days  = flt(getattr(doc, 'total_working_days', None) or
                               getattr(doc, 'total_days_in_month', None) or 0)
            absent_days = flt(getattr(doc, 'absent_days', None) or
                               getattr(doc, 'leave_without_pay', None) or 0)
            paid_days   = flt(getattr(doc, 'payment_days', None) or 0)

            # ── Net Salary from Salary Slip ───────────────────────────────
            ss_net = flt(doc.net_salary)

            # ── Net Salary from Salary Structure Assignment ───────────────
            ssa_net = 0.0
            try:
                ssa_name = frappe.db.get_value(
                    "Salary Structure Assignment",
                    {"employee": doc.employee, "docstatus": 1},
                    "name",
                    order_by="from_date desc"
                )
                if ssa_name:
                    ssa_doc = frappe.get_doc("Salary Structure Assignment", ssa_name)
                    ssa_earnings   = sum(flt(r.amount) for r in (ssa_doc.earnings   or []))
                    ssa_deductions = sum(flt(r.amount) for r in (ssa_doc.deductions or []))
                    ssa_net = ssa_earnings - ssa_deductions
            except Exception:
                frappe.log_error(frappe.get_traceback(), f"SSA net fetch failed for {name}")

            # ── Fetch Additional Salary & Deduction components ────────────
            slip_date = getattr(doc, 'start_date', None)
            additional_salary_components    = {}
            additional_deduction_components = {}

            if slip_date:
                year = str(slip_date.year)
                month_names = [
                    "January","February","March","April","May","June",
                    "July","August","September","October","November","December"
                ]
                month = month_names[slip_date.month - 1]

                for rec in frappe.db.get_all(
                    "Additional Salary",
                    filters={"employee": doc.employee, "year": year, "month": month, "docstatus": 1},
                    fields=["name"]
                ):
                    add_doc = frappe.get_doc("Additional Salary", rec.name)
                    for row in add_doc.components or []:
                        additional_salary_components[row.component_type] = flt(row.amount)

                for rec in frappe.db.get_all(
                    "Additional Deductions",
                    filters={"employee": doc.employee, "year": year, "month": month, "docstatus": 1},
                    fields=["name"]
                ):
                    add_doc = frappe.get_doc("Additional Deductions", rec.name)
                    for row in add_doc.deductions or []:
                        additional_deduction_components[row.component_type] = flt(row.amount)

            # ── Classify deductions from salary slip ──────────────────────
            loan_total          = 0.0
            retention_total     = 0.0
            loan_breakdown      = {}
            loan_names_in_slip  = []
            add_sal_deducted    = {}
            add_ded_deducted    = {}

            for row in doc.deductions:
                comp       = row.salary_component or ""
                amt        = flt(row.amount)
                comp_lower = comp.lower()

                if comp_lower == "loan" or comp_lower.startswith("loan-"):
                    loan_total += amt
                    loan_breakdown[comp] = loan_breakdown.get(comp, 0.0) + amt
                    loan_ref = (
                        getattr(row, "loan_name", None) or
                        getattr(row, "loan_id",   None)
                    )
                    if loan_ref and loan_ref not in loan_names_in_slip:
                        loan_names_in_slip.append(loan_ref)

                elif "retention" in comp_lower:
                    retention_total += amt

                elif comp in additional_salary_components:
                    add_sal_deducted[comp] = amt

                elif comp in additional_deduction_components:
                    add_ded_deducted[comp] = amt

            # ── Scan earnings rows for additional salary components (e.g. OT) ──
            for row in doc.earnings:
                comp = row.salary_component or ""
                amt  = flt(row.amount)
                if comp in additional_salary_components and comp not in add_sal_deducted:
                    add_sal_deducted[comp] = amt

            # ── Fill in 0-amount entries so they still appear in the table ──
            for comp in additional_salary_components:
                if comp not in add_sal_deducted:
                    add_sal_deducted[comp] = 0

            for comp in additional_deduction_components:
                if comp not in add_ded_deducted:
                    add_ded_deducted[comp] = 0

            # ── Loan deferral flag from slip deduction row (masters removed) ─
            loan_deferred = False
            for row in doc.deductions:
                comp_lower = (row.salary_component or "").lower()
                if (comp_lower == "loan" or comp_lower.startswith("loan-")) and int(getattr(row, "is_deferred", 0) or 0):
                    loan_deferred = True
                    break

            # ── Build result ──────────────────────────────────────────────
            result[name] = {
                "employee":      doc.employee,
                "employee_name": doc.employee_name,
                "ssa_net":       round(ssa_net, 2),
                "ss_net":        round(ss_net,  2),
                "total_days":    total_days,
                "absent_days":   absent_days,
                "paid_days":     paid_days,
                "loan": {
                    "total":     round(loan_total, 2),
                    "deferred":  loan_deferred,
                    "breakdown": loan_breakdown,
                },
                "retention":             round(retention_total, 2),
                "additional_salary":     add_sal_deducted,
                "additional_deductions": add_ded_deducted,
            }

        except Exception:
            frappe.log_error(
                frappe.get_traceback(),
                f"get_all_slip_deductions failed for {name}"
            )

    return result


# ── SSA Export ────────────────────────────────────────────────────────────────
@frappe.whitelist()
def export_ssa_to_excel(status_filter=None):
    import openpyxl
    import io
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # ── Determine permitted employees ─────────────────────────────────────────
    permitted_employees = _get_permitted_employees()

    # Build filters
    filters = {"docstatus": ["!=", 2]}
    if status_filter is not None and str(status_filter) != "":
        filters = {"docstatus": int(status_filter)}

    # Add employee restriction if applicable
    if permitted_employees is not None:
        if len(permitted_employees) == 0:
            frappe.throw("You do not have permission to export any records.")
        filters["employee"] = ["in", permitted_employees]

    # Fetch SSA records
    assignments = frappe.get_all(
        "Salary Structure Assignment",
        filters=filters,
        fields=[
            "name", "employee", "employee_name", "company",
            "from_date", "to_date", "salary_structure",
            "designation", "branch", "department", "currency",
            "gross_salary", "total_deductions", "total_employer_contribution",
            "monthly_ctc", "net_salary", "annual_ctc",
            "docstatus", "owner", "amended_from"
        ],
        order_by="from_date desc"
    )

    if not assignments:
        frappe.throw("No Salary Structure Assignments found.")

    ssa_names = [a["name"] for a in assignments]

    # Fetch child rows in bulk
    earnings_rows = frappe.get_all(
        "Salary Details",
        filters={"parent": ["in", ssa_names], "parentfield": "earnings"},
        fields=["parent", "salary_component", "amount"],
        order_by="idx asc"
    )
    deduction_rows = frappe.get_all(
        "Salary Details",
        filters={"parent": ["in", ssa_names], "parentfield": "deductions"},
        fields=["parent", "salary_component", "amount", "employer_contribution"],
        order_by="idx asc"
    )

    # Build lookup maps
    earn_map, ded_map = {}, {}
    all_earn_components, all_ded_components = [], []

    for r in earnings_rows:
        earn_map.setdefault(r.parent, {})[r.salary_component] = r.amount
        if r.salary_component not in all_earn_components:
            all_earn_components.append(r.salary_component)

    for r in deduction_rows:
        ded_map.setdefault(r.parent, {})[r.salary_component] = r.amount
        if r.salary_component not in all_ded_components:
            all_ded_components.append(r.salary_component)

    # Headers
    base_headers = ["ID", "Employee", "Employee Name", "Company",
                    "From Date", "To Date", "Salary Structure",
                    "Designation", "Branch", "Department", "Currency",
                    "Status", "Created By", "Amended From"]
    earn_headers = [f"Earn: {c}" for c in all_earn_components]
    ded_headers  = [f"Ded: {c}"  for c in all_ded_components]
    calc_headers = ["Gross Salary", "Total Deductions", "Employer Contribution",
                    "Monthly CTC", "Net Salary", "Annual CTC"]
    all_headers  = base_headers + earn_headers + ded_headers + calc_headers

    STATUS_MAP = {0: "Draft", 1: "Submitted", 2: "Cancelled"}

    # Data rows
    rows = []
    for a in assignments:
        base = [
            a.name, a.employee, a.employee_name, a.company,
            str(a.from_date) if a.from_date else "",
            str(a.to_date)   if a.to_date   else "",
            a.salary_structure,
            a.designation or "", a.branch or "", a.department or "",
            a.currency or "INR",
            STATUS_MAP.get(a.docstatus, ""),
            a.owner or "", a.amended_from or ""
        ]
        earn_vals = [earn_map.get(a.name, {}).get(c, "") for c in all_earn_components]
        ded_vals  = [ded_map.get(a.name,  {}).get(c, "") for c in all_ded_components]
        calc      = [a.gross_salary or 0, a.total_deductions or 0,
                     a.total_employer_contribution or 0, a.monthly_ctc or 0,
                     a.net_salary or 0, a.annual_ctc or 0]
        rows.append(base + earn_vals + ded_vals + calc)

    # Build workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SSA Export"

    hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    row_font = Font(name="Arial", size=10)
    center   = Alignment(horizontal="center", vertical="center")
    left     = Alignment(horizontal="left")
    right    = Alignment(horizontal="right")
    thin     = Border(
        left  =Side(style="thin", color="DDDDDD"),
        right =Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD")
    )
    fills = {
        "base": PatternFill("solid", start_color="2D2D2D"),
        "earn": PatternFill("solid", start_color="1A5C32"),
        "ded":  PatternFill("solid", start_color="7B1C1C"),
        "calc": PatternFill("solid", start_color="1A3A5C"),
    }
    alt_fill = PatternFill("solid", start_color="F7F7F7")

    def col_section(idx):
        if idx < len(base_headers):                                         return "base"
        if idx < len(base_headers) + len(earn_headers):                    return "earn"
        if idx < len(base_headers) + len(earn_headers) + len(ded_headers): return "ded"
        return "calc"

    # Row 1 — group labels
    ws.row_dimensions[1].height = 16
    groups = [
        ("BASE INFO",    1,                                                        len(base_headers),                                    "3D3D3D"),
        ("EARNINGS",     len(base_headers)+1,                                      len(base_headers)+len(earn_headers),                  "27803E"),
        ("DEDUCTIONS",   len(base_headers)+len(earn_headers)+1,                    len(base_headers)+len(earn_headers)+len(ded_headers), "A02020"),
        ("CALCULATIONS", len(all_headers)-len(calc_headers)+1,                    len(all_headers),                                     "1E4F80"),
    ]
    for label, sc, ec, color in groups:
        if sc > ec:
            continue
        cell = ws.cell(row=1, column=sc, value=label)
        cell.font      = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        cell.fill      = PatternFill("solid", start_color=color)
        cell.alignment = center
        if sc < ec:
            ws.merge_cells(start_row=1, start_column=sc, end_row=1, end_column=ec)

    # Row 2 — column headers
    ws.row_dimensions[2].height = 26
    for ci, hdr in enumerate(all_headers, start=1):
        cell = ws.cell(row=2, column=ci, value=hdr)
        cell.font      = hdr_font
        cell.fill      = fills[col_section(ci - 1)]
        cell.alignment = center
        cell.border    = thin

    # Numeric column ranges (1-based)
    earn_start = len(base_headers) + 1
    earn_end   = earn_start + len(earn_headers) - 1
    ded_start  = earn_end + 1
    ded_end    = ded_start + len(ded_headers) - 1
    calc_start = ded_end + 1
    calc_end   = len(all_headers)

    # Data rows start at row 3
    for ri, row_data in enumerate(rows, start=3):
        use_alt = (ri % 2 == 0)
        for ci, val in enumerate(row_data, start=1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font   = row_font
            cell.border = thin
            if use_alt:
                cell.fill = alt_fill
            is_num = (earn_start <= ci <= earn_end or
                      ded_start  <= ci <= ded_end  or
                      calc_start <= ci <= calc_end)
            if is_num and val != "":
                cell.number_format = "#,##0.00"
                cell.alignment     = right
            else:
                cell.alignment = left

    # Auto column widths
    for ci, hdr in enumerate(all_headers, start=1):
        max_len = len(str(hdr))
        for ri in range(3, min(3 + len(rows), 103)):
            v = ws.cell(row=ri, column=ci).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 3, 28)

    ws.freeze_panes = "C3"

    # ── Save to bytes buffer ──────────────────────────────────────────────────
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    file_content = buffer.read()

    # ── Delete previous export if exists ─────────────────────────────────────
    existing = frappe.db.get_value(
        "File",
        {"file_name": "ssa_export.xlsx", "is_private": 1},
        "name"
    )
    if existing:
        frappe.delete_doc("File", existing, force=True)

    # ── Save as proper Frappe File doc (private, session-authenticated) ───────
    file_doc = frappe.get_doc({
        "doctype": "File",
        "file_name": "ssa_export.xlsx",
        "is_private": 1,
        "content": file_content,
        "decode": False
    })
    file_doc.save(ignore_permissions=True)

    return {"file_url": file_doc.file_url, "rows": len(rows)}