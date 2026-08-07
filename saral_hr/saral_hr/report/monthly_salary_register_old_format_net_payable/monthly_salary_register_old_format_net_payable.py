# Copyright (c) 2026, sj and Contributors
# See license.txt

"""Monthly Salary Register Old Format Net Payable — Desk + PDF + Excel."""

from __future__ import annotations

import calendar
import io
import json
import re
from datetime import date

import frappe
from frappe import _
from frappe.utils import cint, flt, get_last_day, getdate
from frappe.utils.pdf import get_pdf

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONTH_MAP = {
	"January": 1,
	"February": 2,
	"March": 3,
	"April": 4,
	"May": 5,
	"June": 6,
	"July": 7,
	"August": 8,
	"September": 9,
	"October": 10,
	"November": 11,
	"December": 12,
}

DAY_COLS = [
	("day_p", "P", "present_days"),
	("day_r", "R", "weekly_offs_taken"),
	("day_h", "H", "holidays_taken"),
	("day_c", "C", "total_casual_leaves"),
	("day_t", "T", "total_on_tour"),
	("day_y", "Y", "total_comp_off"),
	("day_od", "OD", None),  # On Duty — not in product yet
	("day_e", "E", "total_earned_leaves"),
	("day_i", "I", None),  # ESIC Leave — not in product yet
	("day_a", "A", "absent_days"),
]

IDENTITY_FNS = {"sr_no", "employee_name", "bank_name", "ifsc", "bank_account"}
COMPRESSED_FORMAT = "Compressed"
FULL_FORMAT = "Full"

ACTUAL_HDR = "#F4A460"
EARN_HDR = "#90EE90"
DED_HDR = "#F08080"
SIG = (
	'<table style="width:100%;margin-top:36px;border-collapse:collapse;table-layout:fixed;">'
	"<tr>"
	+ "".join(
		f'<td style="width:33.33%;text-align:center;vertical-align:top;padding:0 12px;">'
		f'<div style="border-top:1px solid #000;margin:0 auto 6px auto;width:85%;"></div>'
		f'<div style="font-size:10px;">{lbl}</div></td>'
		for lbl in ("Prepared By", "Checked By", "Authorised Signatory")
	)
	+ "</tr></table>"
)


def execute(filters=None):
	payload = build_register(filters or {})
	return payload["columns"], payload["data"]


def build_register(filters):
	"""Build columns + data + meta for Desk / PDF / Excel."""
	f = _normalize_filters(filters)
	compressed = f.format == COMPRESSED_FORMAT
	base_meta = {
		"company": f.company,
		"month": f.month,
		"year": f.year,
		"format": f.format,
		"actual_comps": [],
		"earn_comps": [],
		"ded_comps": [],
	}

	if not f.company or not f.year or not f.month:
		return {
			"columns": _compressed_columns() if compressed else _empty_columns(),
			"data": [],
			"meta": base_meta,
		}

	start = date(int(f.year), MONTH_MAP[f.month], 1)
	end = getdate(get_last_day(start))
	month_days = calendar.monthrange(int(f.year), MONTH_MAP[f.month])[1]
	base_meta.update({"start": str(start), "end": str(end), "month_days": month_days})

	slips = _load_slips(f.company, start, end, f.population)
	if not slips:
		return {
			"columns": _compressed_columns() if compressed else _fixed_columns([], [], []),
			"data": [],
			"meta": base_meta,
		}

	cl_ids = [s.employee for s in slips]
	cl_map = _company_link_map(cl_ids)
	emp_ids = [cl_map[c]["employee"] for c in cl_ids if cl_map.get(c) and cl_map[c].get("employee")]
	emp_map = _employee_bank_map(emp_ids)

	ssa_by_cl = {}
	for s in slips:
		ssa_by_cl[s.employee] = _resolve_ssa(s.employee, start, end)

	if compressed:
		return _build_compressed(slips, cl_map, emp_map, ssa_by_cl, month_days, base_meta)

	actual_union = []
	seen_act = set()
	for s in slips:
		for comp, amt in _ssa_earnings(ssa_by_cl.get(s.employee)):
			if comp not in seen_act:
				seen_act.add(comp)
				actual_union.append(comp)

	earn_maps, ded_maps = _slip_component_maps([s.name for s in slips])
	earn_union = list(actual_union)
	seen_earn = set(actual_union)
	for slip_name in earn_maps:
		for comp in earn_maps[slip_name]:
			if comp not in seen_earn:
				seen_earn.add(comp)
				earn_union.append(comp)

	ded_union = []
	seen_ded = set()
	for slip_name in ded_maps:
		for comp in ded_maps[slip_name]:
			if comp not in seen_ded:
				seen_ded.add(comp)
				ded_union.append(comp)

	columns = _fixed_columns(actual_union, earn_union, ded_union)
	rows = []
	for idx, s in enumerate(slips, start=1):
		cl = cl_map.get(s.employee) or {}
		emp = emp_map.get(cl.get("employee")) or {}
		row = {
			"sr_no": idx,
			"employee_name": s.employee_name or cl.get("full_name") or s.employee,
			"bank_name": emp.get("bank_name") or "",
			"ifsc": emp.get("ifsc_code") or "",
			"bank_account": emp.get("account_number") or "",
			"total_month_day": month_days,
			"total_paid_days": flt(s.payment_days),
			"total_gross": 0.0,
			"total_earning": flt(s.total_earnings),
			"total_deductions": flt(s.total_deductions),
			"net_payable": flt(s.net_salary),
			"_is_total": 0,
		}
		for fn, _lbl, field in DAY_COLS:
			row[fn] = flt(getattr(s, field, 0) or 0) if field else None

		ssa_amts = dict(_ssa_earnings(ssa_by_cl.get(s.employee)))
		gross = 0.0
		for comp in actual_union:
			fn = _act_fn(comp)
			val = flt(ssa_amts.get(comp))
			row[fn] = val
			gross += val
		row["total_gross"] = flt(gross, 2)

		slip_earn = earn_maps.get(s.name) or {}
		for comp in earn_union:
			row[_earn_fn(comp)] = flt(slip_earn.get(comp))

		slip_ded = ded_maps.get(s.name) or {}
		for comp in ded_union:
			row[_ded_fn(comp)] = flt(slip_ded.get(comp))

		rows.append(row)

	if rows:
		rows.append(_totals_row(rows, actual_union, earn_union, ded_union, month_days))

	base_meta.update({
		"actual_comps": actual_union,
		"earn_comps": earn_union,
		"ded_comps": ded_union,
	})
	return {"columns": columns, "data": rows, "meta": base_meta}


def _build_compressed(slips, cl_map, emp_map, ssa_by_cl, month_days, base_meta):
	rows = []
	for idx, s in enumerate(slips, start=1):
		cl = cl_map.get(s.employee) or {}
		emp = emp_map.get(cl.get("employee")) or {}
		ssa_amts = dict(_ssa_earnings(ssa_by_cl.get(s.employee)))
		gross = flt(sum(flt(v) for v in ssa_amts.values()), 2)
		rows.append({
			"sr_no": idx,
			"employee_name": s.employee_name or cl.get("full_name") or s.employee,
			"ifsc": emp.get("ifsc_code") or "",
			"bank_account": emp.get("account_number") or "",
			"total_month_day": month_days,
			"total_paid_days": flt(s.payment_days),
			"total_gross": gross,
			"total_earning": flt(s.total_earnings),
			"total_deductions": flt(s.total_deductions),
			"net_payable": flt(s.net_salary),
			"_is_total": 0,
		})

	if rows:
		rows.append(_compressed_totals_row(rows))

	return {
		"columns": _compressed_columns(),
		"data": rows,
		"meta": base_meta,
	}


# ─── Filters / loaders ───────────────────────────────────────────────────────


def _normalize_filters(filters):
	if isinstance(filters, str):
		filters = json.loads(filters)
	f = frappe._dict(filters or {})
	pop = (f.get("population") or "All").strip()
	if pop.lower() == "all":
		pop = "All"
	f.population = pop
	fmt = (f.get("format") or FULL_FORMAT).strip()
	f.format = COMPRESSED_FORMAT if fmt == COMPRESSED_FORMAT else FULL_FORMAT
	return f


def _load_slips(company, start, end, population):
	params = {
		"company": company,
		"start": str(start),
		"end": str(end),
	}
	cat_clause = ""
	if population and population != "All":
		cat_clause = "AND cl.category = %(category)s"
		params["category"] = population

	return frappe.db.sql(
		f"""
		SELECT
			ss.name, ss.employee, ss.employee_name,
			ss.payment_days, ss.present_days, ss.weekly_offs_taken, ss.holidays_taken,
			ss.total_casual_leaves, ss.total_on_tour, ss.total_comp_off,
			ss.total_earned_leaves, ss.absent_days,
			ss.total_earnings, ss.total_deductions, ss.net_salary
		FROM `tabSalary Slip` ss
		INNER JOIN `tabCompany Link` cl ON cl.name = ss.employee
		WHERE ss.docstatus = 1
		  AND ss.company = %(company)s
		  AND ss.start_date = %(start)s
		  AND ss.end_date = %(end)s
		  {cat_clause}
		ORDER BY IFNULL(NULLIF(cl.employee, ''), ss.employee) ASC
		""",
		params,
		as_dict=True,
	)


def _company_link_map(names):
	if not names:
		return {}
	rows = frappe.db.get_all(
		"Company Link",
		filters={"name": ["in", names]},
		fields=["name", "employee", "full_name", "category"],
	)
	return {r.name: r for r in rows}


def _employee_bank_map(names):
	if not names:
		return {}
	rows = frappe.db.get_all(
		"Employee",
		filters={"name": ["in", names]},
		fields=["name", "bank_name", "ifsc_code", "account_number"],
	)
	return {r.name: r for r in rows}


def _resolve_ssa(employee, start, end):
	rows = frappe.db.sql(
		"""
		SELECT name, from_date, to_date
		FROM `tabSalary Structure Assignment`
		WHERE employee = %(employee)s
		  AND docstatus = 1
		  AND from_date <= %(start)s
		  AND (to_date IS NULL OR to_date >= %(end)s)
		ORDER BY from_date DESC
		LIMIT 1
		""",
		{"employee": employee, "start": str(start), "end": str(end)},
		as_dict=True,
	)
	return rows[0].name if rows else None


def _ssa_earnings(ssa_name):
	if not ssa_name:
		return []
	return frappe.db.sql(
		"""
		SELECT salary_component, amount
		FROM `tabSalary Details`
		WHERE parent = %(parent)s
		  AND parenttype = 'Salary Structure Assignment'
		  AND parentfield = 'earnings'
		  AND IFNULL(amount, 0) != 0
		ORDER BY idx ASC
		""",
		{"parent": ssa_name},
	)


def _slip_component_maps(slip_names):
	earn = {n: {} for n in slip_names}
	ded = {n: {} for n in slip_names}
	if not slip_names:
		return earn, ded

	rows = frappe.db.sql(
		"""
		SELECT sd.parent, sd.parentfield, sd.salary_component, sd.amount,
		       IFNULL(sc.employer_contribution, 0) AS employer_contribution
		FROM `tabSalary Details` sd
		LEFT JOIN `tabSalary Component` sc ON sc.name = sd.salary_component
		WHERE sd.parent IN %(parents)s
		  AND sd.parenttype = 'Salary Slip'
		  AND sd.parentfield IN ('earnings', 'deductions')
		ORDER BY sd.idx ASC
		""",
		{"parents": slip_names},
		as_dict=True,
	)
	for r in rows:
		amt = flt(r.amount)
		if not amt:
			continue
		if r.parentfield == "earnings":
			earn[r.parent][r.salary_component] = amt
		elif r.parentfield == "deductions" and not int(r.employer_contribution or 0):
			ded[r.parent][r.salary_component] = amt
	return earn, ded


# ─── Columns / rows ──────────────────────────────────────────────────────────


def _slug(comp):
	s = re.sub(r"[^A-Za-z0-9]+", "_", comp or "").strip("_").lower()
	return s or "comp"


def _act_fn(comp):
	return f"act_{_slug(comp)}"


def _earn_fn(comp):
	return f"earn_{_slug(comp)}"


def _ded_fn(comp):
	return f"ded_{_slug(comp)}"


def _col(label, fn, ft="Data", width=100, **kw):
	d = {"label": _(label), "fieldname": fn, "fieldtype": ft, "width": width}
	d.update(kw)
	return d


def _empty_columns():
	return _fixed_columns([], [], [])


def _compressed_columns():
	# Desk already shows DataTable serial — no SR. NO. column.
	# <br> = two-line headers (Desk / PDF); Excel swaps to \\n.
	return [
		_col("Full<br>Name", "employee_name", width=180),
		_col("IFSC<br>Code", "ifsc", width=140, align="left"),
		_col("Account<br>Number", "bank_account", width=165, align="left"),
		_col("Days in<br>Month", "total_month_day", "Int", 90),
		_col("Paid For<br>Days", "total_paid_days", "Float", 95, precision=1),
		_col("Actual<br>Gross", "total_gross", "Currency", 120),
		_col("Earning<br>Gross", "total_earning", "Currency", 120),
		_col("Deductions", "total_deductions", "Currency", 115),
		_col("Net<br>Payable", "net_payable", "Currency", 120),
	]


def _fmt_paid_days(val):
	"""Whole days without .0; keep one decimal only when fractional."""
	if val is None or val == "":
		return ""
	v = flt(val)
	if abs(v - round(v)) < 1e-9:
		return str(int(round(v)))
	return f"{v:.1f}"


def _columns_for_export(payload):
	"""PDF/Excel keep an SR. NO. (Desk uses DataTable serial instead)."""
	cols = list(payload["columns"])
	if payload.get("meta", {}).get("format") == COMPRESSED_FORMAT:
		cols = [_col("SR.<br>NO.", "sr_no", "Int", 45)] + cols
	return cols


def _colgroup_html(cols):
	"""Proportional <col> widths so PDF does not force equal columns."""
	total = sum(cint(c.get("width") or 100) for c in cols) or 1
	parts = []
	for c in cols:
		pct = 100.0 * cint(c.get("width") or 100) / total
		parts.append(f'<col style="width:{pct:.2f}%">')
	return "<colgroup>" + "".join(parts) + "</colgroup>"


def _excel_col_width(col):
	"""Map Desk px-ish widths to Excel character widths (breathing room)."""
	fn = col.get("fieldname")
	px = cint(col.get("width") or 100)
	# Excel width ≈ characters; tune bank fields so IFSC / account don't clip
	if fn == "sr_no":
		return 6
	if fn == "employee_name":
		return max(18, int(px / 9))
	if fn == "ifsc":
		return max(14, int(px / 9))
	if fn == "bank_account":
		return max(16, int(px / 9))
	if fn in ("total_gross", "total_earning", "total_deductions", "net_payable"):
		return max(12, int(px / 9))
	return max(10, int(px / 9))


def _fixed_columns(actual_comps, earn_comps, ded_comps):
	cols = [
		_col("SR. NO.", "sr_no", "Int", 50),
		_col("Full Name Of The Employee", "employee_name", width=180),
		_col("BANK", "bank_name", width=90),
		_col("IFSC", "ifsc", width=130, align="left"),
		_col("BANK ACCOUNT NO", "bank_account", width=140, align="left"),
		_col("TOTAL MONTH DAY", "total_month_day", "Int", 70),
	]
	for fn, lbl, _ in DAY_COLS:
		cols.append(_col(lbl, fn, "Float", 45, precision=1))
	cols.append(_col("TOTAL PAID DAYS", "total_paid_days", "Float", 80, precision=1))

	for comp in actual_comps:
		cols.append(_col(comp, _act_fn(comp), "Currency", 90))
	cols.append(_col("Total Gross", "total_gross", "Currency", 100))

	for comp in earn_comps:
		cols.append(_col(comp, _earn_fn(comp), "Currency", 90))
	cols.append(_col("Total Earning", "total_earning", "Currency", 100))

	for comp in ded_comps:
		cols.append(_col(comp, _ded_fn(comp), "Currency", 90))
	cols.append(_col("Total Deductions", "total_deductions", "Currency", 110))
	cols.append(_col("Net payable", "net_payable", "Currency", 110))
	return cols


def _totals_row(rows, actual_comps, earn_comps, ded_comps, month_days):
	tot = {
		"sr_no": "",
		"employee_name": "Total",
		"bank_name": "",
		"ifsc": "",
		"bank_account": "",
		"total_month_day": "",
		"_is_total": 1,
	}
	sum_fns = ["total_paid_days", "total_gross", "total_earning", "total_deductions", "net_payable"]
	for fn, _, field in DAY_COLS:
		if field:
			sum_fns.append(fn)
	for comp in actual_comps:
		sum_fns.append(_act_fn(comp))
	for comp in earn_comps:
		sum_fns.append(_earn_fn(comp))
	for comp in ded_comps:
		sum_fns.append(_ded_fn(comp))

	for fn in sum_fns:
		tot[fn] = flt(sum(flt(r.get(fn)) for r in rows), 2)
	for fn, _, field in DAY_COLS:
		if not field:
			tot[fn] = None
	return tot


def _compressed_totals_row(rows):
	tot = {
		"sr_no": "",
		"employee_name": "Total",
		"ifsc": "",
		"bank_account": "",
		"total_month_day": "",
		"total_paid_days": "",
		"_is_total": 1,
	}
	for fn in (
		"total_gross",
		"total_earning",
		"total_deductions",
		"net_payable",
	):
		tot[fn] = flt(sum(flt(r.get(fn)) for r in rows), 2)
	return tot


# ─── PDF ─────────────────────────────────────────────────────────────────────


@frappe.whitelist()
def print_report(filters):
	payload = build_register(filters)
	html = _build_html(payload)
	return _save_pdf(html, "Monthly_Salary_Register")


def _build_html(payload):
	cols = _columns_for_export(payload)
	data = payload["data"]
	meta = payload["meta"]
	company = meta.get("company") or ""
	start = getdate(meta.get("start")) if meta.get("start") else None
	end = getdate(meta.get("end")) if meta.get("end") else None
	period = ""
	if start and end:
		period = f"{start.strftime('%d-%b-%Y')} to {end.strftime('%d-%b-%Y')}"

	compressed = meta.get("format") == COMPRESSED_FORMAT
	actual_set = {_act_fn(c) for c in meta.get("actual_comps") or []}
	earn_set = {_earn_fn(c) for c in meta.get("earn_comps") or []}
	ded_set = {_ded_fn(c) for c in meta.get("ded_comps") or []}

	def hdr_bg(fn):
		if compressed:
			return "#f0f0f0"
		if fn in actual_set or fn == "total_gross":
			return ACTUAL_HDR
		if fn in earn_set or fn == "total_earning":
			return EARN_HDR
		if fn in ded_set or fn == "total_deductions":
			return DED_HDR
		return "#f0f0f0"

	th = "".join(
		f'<th style="padding:4px 6px 4px 8px;font-size:6.5px;'
		f'background:{hdr_bg(c["fieldname"])};vertical-align:middle;">'
		f'{c["label"]}</th>'
		for c in cols
	)

	body = ""
	data_idx = 0
	for row in data:
		tds = ""
		is_tot = row.get("_is_total")
		if not is_tot:
			data_idx += 1
		for c in cols:
			fn = c["fieldname"]
			val = row.get(fn)
			if val is None or val == "":
				disp = ""
			elif fn == "total_paid_days":
				disp = "" if is_tot else _fmt_paid_days(val)
			elif isinstance(val, float):
				disp = f"{val:,.2f}"
			else:
				disp = str(val)
			if is_tot:
				bg = "#e8e8e8"
			elif data_idx % 2 == 0:
				bg = "#f5f5f5"
			else:
				bg = "#fff"
			align = "left" if fn in ("employee_name", "bank_name", "ifsc", "bank_account") else "right"
			if fn == "sr_no":
				align = "center"
			weight = "bold" if is_tot else "normal"
			pad = "2px 4px 2px 8px"
			row_cls = "tot" if is_tot else "data"
			tds += (
				f'<td class="{row_cls}" style="padding:{pad};font-size:6.5px;'
				f'text-align:{align};background:{bg};font-weight:{weight};'
				f'white-space:nowrap;overflow:hidden;">{disp}</td>'
			)
		body += f"<tr>{tds}</tr>"

	colgroup = _colgroup_html(cols)
	return f"""<!DOCTYPE html><html><head><meta charset="UTF-8">
<style>
*{{margin:0;padding:0;box-sizing:border-box}}
body{{font-family:Arial,sans-serif;font-size:7px;color:#000}}
.hdr{{text-align:center;margin-bottom:6px}}
.hdr .co{{font-size:14px;font-weight:900}}
.hdr .ttl{{font-size:12px;font-weight:700;margin-top:2px}}
.hdr .per{{font-size:10px;margin-top:2px}}
table{{width:100%;border-collapse:collapse;table-layout:fixed}}
th,td{{word-wrap:break-word}}
/* Vertical lines stay solid; horizontal lines in data area are thin dashed */
th{{
  border-left:1px solid #000;
  border-right:1px solid #000;
  border-top:1px solid #000;
  border-bottom:1px solid #000;
}}
td.data{{
  border-left:1px solid #000;
  border-right:1px solid #000;
  border-top:none;
  border-bottom:0.5px dashed #999;
}}
td.tot{{
  border-left:1px solid #000;
  border-right:1px solid #000;
  border-top:1px solid #000;
  border-bottom:1px solid #000;
}}
</style></head><body>
<div class="hdr">
  <div class="co">{frappe.utils.escape_html(company)}</div>
  <div class="ttl">Employee Salary Sheet</div>
  <div class="per">For the Period {period} &nbsp;|&nbsp; Month Days: {meta.get("month_days") or ""}</div>
</div>
<table>{colgroup}<thead><tr>{th}</tr></thead><tbody>{body}</tbody></table>
{SIG}
</body></html>"""


def _save_pdf(html, prefix):
	pdf = get_pdf(
		html,
		options={
			"page-size": "A4",
			"orientation": "Landscape",
			"margin-top": "6mm",
			"margin-right": "4mm",
			"margin-bottom": "6mm",
			# Extra left margin for filing / hole-punch
			"margin-left": "18mm",
			"encoding": "UTF-8",
			"no-outline": None,
		},
	)
	ts = frappe.utils.now_datetime().strftime("%Y%m%d_%H%M%S")
	fn = f"{prefix}_{ts}.pdf"
	with open(frappe.utils.get_files_path(fn, is_private=0), "wb") as fh:
		fh.write(pdf)
	doc = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": fn,
			"is_private": 0,
			"file_url": f"/files/{fn}",
		}
	)
	doc.insert(ignore_permissions=True)
	frappe.db.commit()
	return doc.file_url


# ─── Excel ───────────────────────────────────────────────────────────────────


@frappe.whitelist()
def excel_report(filters):
	payload = build_register(filters)
	wb = _build_workbook(payload)
	return _save_excel(wb, "Monthly_Salary_Register")


def _build_workbook(payload):
	cols = _columns_for_export(payload)
	data = payload["data"]
	meta = payload["meta"]
	wb = Workbook()
	ws = wb.active
	ws.title = "Salary Register"

	thin = Border(
		left=Side(style="thin", color="000000"),
		right=Side(style="thin", color="000000"),
		top=Side(style="thin", color="000000"),
		bottom=Side(style="thin", color="000000"),
	)
	company = meta.get("company") or ""
	start = getdate(meta.get("start")) if meta.get("start") else None
	end = getdate(meta.get("end")) if meta.get("end") else None
	period = ""
	if start and end:
		period = f"{start.strftime('%d-%b-%Y')} to {end.strftime('%d-%b-%Y')}"

	ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(cols), 1))
	c1 = ws.cell(1, 1, f"{company} - Employee Salary Sheet")
	c1.font = Font(name="Arial", size=14, bold=True)
	c1.alignment = Alignment(horizontal="center")

	ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(cols), 1))
	c2 = ws.cell(2, 1, f"For the Period {period}  |  Month Days: {meta.get('month_days') or ''}")
	c2.font = Font(name="Arial", size=10)
	c2.alignment = Alignment(horizontal="center")

	compressed = meta.get("format") == COMPRESSED_FORMAT
	actual_set = {_act_fn(c) for c in meta.get("actual_comps") or []}
	earn_set = {_earn_fn(c) for c in meta.get("earn_comps") or []}
	ded_set = {_ded_fn(c) for c in meta.get("ded_comps") or []}

	def fill_for(fn):
		if compressed:
			return PatternFill("solid", fgColor="F0F0F0")
		if fn in actual_set or fn == "total_gross":
			return PatternFill("solid", fgColor="F4A460")
		if fn in earn_set or fn == "total_earning":
			return PatternFill("solid", fgColor="90EE90")
		if fn in ded_set or fn == "total_deductions":
			return PatternFill("solid", fgColor="F08080")
		return PatternFill("solid", fgColor="F0F0F0")

	hdr_row = 4
	for i, col in enumerate(cols, start=1):
		cell = ws.cell(hdr_row, i, (col["label"] or "").replace("<br>", "\n"))
		cell.font = Font(name="Arial", size=8, bold=True)
		cell.fill = fill_for(col["fieldname"])
		cell.border = thin
		cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")

	tot_fill = PatternFill("solid", fgColor="E8E8E8")

	for r_idx, row in enumerate(data, start=hdr_row + 1):
		is_tot = row.get("_is_total")
		for c_idx, col in enumerate(cols, start=1):
			fn = col["fieldname"]
			val = row.get(fn)
			if val is None:
				val = ""
			cell = ws.cell(r_idx, c_idx, val if val != "" else None)
			cell.font = Font(name="Arial", size=8, bold=bool(is_tot))
			cell.border = thin
			if is_tot:
				cell.fill = tot_fill
			if is_tot and fn == "total_paid_days":
				cell.value = None
			elif fn == "bank_account" or fn == "ifsc":
				cell.alignment = Alignment(horizontal="left")
			elif fn == "total_paid_days" and isinstance(val, (int, float)):
				v = flt(val)
				if abs(v - round(v)) < 1e-9:
					cell.value = int(round(v))
					cell.number_format = "0"
				else:
					cell.number_format = "0.0"
				cell.alignment = Alignment(horizontal="right")
			elif col.get("fieldtype") in ("Currency", "Float") and isinstance(val, (int, float)):
				cell.number_format = "#,##0.00"
				cell.alignment = Alignment(horizontal="right")

	# Signatures — one row, equally spaced across the sheet
	sig_row = hdr_row + len(data) + 3
	labels = ("Prepared By", "Checked By", "Authorised Signatory")
	n_cols = max(len(cols), 1)
	# place each label at the center of its third of the columns
	for i, lbl in enumerate(labels):
		start = int(i * n_cols / 3) + 1
		end = int((i + 1) * n_cols / 3)
		if end < start:
			end = start
		mid = (start + end) // 2
		if start != end:
			ws.merge_cells(start_row=sig_row, start_column=start, end_row=sig_row, end_column=end)
			ws.merge_cells(
				start_row=sig_row + 1, start_column=start, end_row=sig_row + 1, end_column=end
			)
		line = ws.cell(sig_row, start, "________________")
		line.alignment = Alignment(horizontal="center")
		line.font = Font(name="Arial", size=9)
		lab = ws.cell(sig_row + 1, start, lbl)
		lab.alignment = Alignment(horizontal="center")
		lab.font = Font(name="Arial", size=9)

	for i, col in enumerate(cols, start=1):
		ws.column_dimensions[get_column_letter(i)].width = _excel_col_width(col)
	ws.row_dimensions[hdr_row].height = 32
	return wb


def _save_excel(wb, prefix):
	buf = io.BytesIO()
	wb.save(buf)
	content = buf.getvalue()
	ts = frappe.utils.now_datetime().strftime("%Y%m%d_%H%M%S")
	fn = f"{prefix}_{ts}.xlsx"
	path = frappe.utils.get_files_path(fn, is_private=0)
	with open(path, "wb") as fh:
		fh.write(content)
	doc = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": fn,
			"is_private": 0,
			"file_url": f"/files/{fn}",
		}
	)
	doc.insert(ignore_permissions=True)
	frappe.db.commit()
	return doc.file_url
