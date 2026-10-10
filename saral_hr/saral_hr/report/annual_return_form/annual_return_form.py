# Copyright (c) 2026, sj and Contributors
# See license.txt

"""Annual Return Form — one workbook with the five sheets from the factory return."""

from __future__ import annotations

import calendar
import io
import json
from collections import Counter
from datetime import date, datetime

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MONTHS = (
	"January",
	"February",
	"March",
	"April",
	"May",
	"June",
	"July",
	"August",
	"September",
	"October",
	"November",
	"December",
)
FULL_DAY = {"Present", "On Tour", "Earned Comp Off"}
LEAVE_DAY = {"Earned Leave", "Casual Leave", "Comp Off"}
CLOSED_DAY = {"Weekly Off", "Holiday"}
HOURS_PER_DAY = 8
ANNUAL_LEAVE_DAYS = 240
WEEKDAYS = {
	"Monday": 0,
	"Tuesday": 1,
	"Wednesday": 2,
	"Thursday": 3,
	"Friday": 4,
	"Saturday": 5,
	"Sunday": 6,
}
BANDS = (
	("adult", "Adult Workers"),
	("adolescent", "Worker of age (15 yrs < 18 yrs)"),
	("child", "Child Labor (14 yrs to 15 yrs)"),
)
SEXES = (("male", "Male"), ("female", "Female"))

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SUB_FILL = PatternFill("solid", fgColor="D6DCE4")
TOTAL_FILL = PatternFill("solid", fgColor="E8E8E8")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E79")
HEADER_FONT = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
SUB_FONT = Font(name="Calibri", size=9, bold=True)
BODY_FONT = Font(name="Calibri", size=9)
BOLD_FONT = Font(name="Calibri", size=9, bold=True)
THIN = Border(
	left=Side(style="thin", color="000000"),
	right=Side(style="thin", color="000000"),
	top=Side(style="thin", color="000000"),
	bottom=Side(style="thin", color="000000"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def execute(filters=None):
	filters = _normalize_filters(filters)
	year = _year(filters)
	if not filters.get("company") or not year:
		return _working_columns(year or date.today().year), []
	pack = _load(filters.company, year)
	return _working_columns(year), _working_rows(pack)


@frappe.whitelist()
def excel_report(filters):
	filters = _normalize_filters(filters)
	year = _year(filters)
	if not filters.get("company") or not year:
		frappe.throw(_("Please select Company and Year."))
	pack = _load(filters.company, year)
	if not pack["working"]:
		frappe.throw(_("No employees found for this company and year."))
	wb = _workbook(pack, filters.company, year)
	return _save_excel(wb, "Annual_Return_Form")


def worked_weight(status):
	if status in FULL_DAY:
		return 1.0
	if status == "Half Day":
		return 0.5
	return 0.0


def on_rolls(joining, left, start, end):
	joining = _as_date(joining)
	left = _as_date(left)
	if not joining or joining > end:
		return False
	if left and left < start:
		return False
	return True


def age_band(dob, on):
	"""Adult is 18+, adolescent is 15 to under 18, child is under 15. Missing birth date counts as adult."""
	dob = _as_date(dob)
	if not dob:
		return "adult"
	years = on.year - dob.year - ((on.month, on.day) < (dob.month, dob.day))
	if years >= 18:
		return "adult"
	if years >= 15:
		return "adolescent"
	return "child"


def is_female(gender):
	return (gender or "").strip().lower() in {"female", "f"}


def eligible_for_annual_leave(joining, worked_days, year):
	"""Factories Act annual leave: 240 days, or two-thirds of the days after a mid-year joining."""
	joining = _as_date(joining)
	start = date(year, 1, 1)
	end = date(year, 12, 31)
	worked = flt(worked_days)
	if not joining or joining <= start:
		return worked + 1e-9 >= ANNUAL_LEAVE_DAYS
	span = (end - joining).days + 1
	return worked + 1e-9 >= (2.0 / 3.0) * span


def is_agriculture(category):
	return "agricult" in (category or "").lower()


def assemble(employees, attendance, daily_rate, year):
	start = date(year, 1, 1)
	end = date(year, 12, 31)
	people = []
	for emp in employees or []:
		if is_agriculture(emp.get("category")):
			continue
		if not on_rolls(emp.get("date_of_joining"), emp.get("left_date"), start, end):
			continue
		people.append(emp)

	by_code = {emp.get("name"): emp for emp in people}
	marks = {}
	for row in list(attendance or []) + list(daily_rate or []):
		code = row.get("employee")
		if code not in by_code:
			continue
		day = _as_date(row.get("attendance_date") or row.get("date"))
		if not day or day.year != year:
			continue
		slot = marks.setdefault(code, {}).setdefault(day, {"weight": 0.0, "status": "", "ot": 0.0, "leave": False})
		weight = worked_weight(row.get("status"))
		if weight > slot["weight"]:
			slot["weight"] = weight
			slot["status"] = row.get("status") or ""
		elif not slot["status"]:
			slot["status"] = row.get("status") or ""
		if row.get("status") in LEAVE_DAY:
			slot["leave"] = True
		slot["ot"] += flt(row.get("ot_hours"))

	weekday = _common_weekly_off(people)
	plant_open = set()
	holiday_dates = set()
	for code, days in marks.items():
		for day, slot in days.items():
			if slot["status"] and slot["status"] not in CLOSED_DAY:
				plant_open.add(day)
			if slot["status"] == "Holiday":
				holiday_dates.add(day)

	weekly_off_dates = set()
	cursor = start
	while cursor <= end:
		if cursor.weekday() == weekday and cursor not in plant_open:
			weekly_off_dates.add(cursor)
		cursor = _next_day(cursor)
	holiday_dates = {day for day in holiday_dates if day not in plant_open and day not in weekly_off_dates}
	year_days = (end - start).days + 1
	operation_days = year_days - len(weekly_off_dates) - len(holiday_dates)

	working = []
	month_mandays = {(m, sex): 0.0 for m in range(1, 13) for sex in ("male", "female")}
	month_ot = {m: 0.0 for m in range(1, 13)}
	month_head = {(m, sex): 0 for m in range(1, 13) for sex in ("male", "female")}
	groups = {
		(band, sex): {"on_duty": 0, "eligible": 0, "sanctioned": 0, "left": 0, "mandays": 0.0, "ot": 0.0}
		for band, _label in BANDS
		for sex, _sex_label in SEXES
	}

	for emp in sorted(people, key=lambda row: (row.get("name") or "")):
		code = emp.get("name")
		sex = "female" if is_female(emp.get("gender")) else "male"
		band = age_band(emp.get("date_of_birth"), end)
		months = [0.0] * 12
		leave_taken = False
		ot_total = 0.0
		for day, slot in marks.get(code, {}).items():
			months[day.month - 1] += slot["weight"]
			month_mandays[(day.month, sex)] += slot["weight"]
			month_ot[day.month] += slot["ot"]
			ot_total += slot["ot"]
			leave_taken = leave_taken or slot["leave"]
		worked = sum(months)
		groups[(band, sex)]["on_duty"] += 1
		groups[(band, sex)]["mandays"] += worked
		groups[(band, sex)]["ot"] += ot_total
		if eligible_for_annual_leave(emp.get("date_of_joining"), worked, year):
			groups[(band, sex)]["eligible"] += 1
		if leave_taken:
			groups[(band, sex)]["sanctioned"] += 1
		left = _as_date(emp.get("left_date"))
		if left and left.year == year:
			groups[(band, sex)]["left"] += 1
		for month in range(1, 13):
			m_start = date(year, month, 1)
			m_end = date(year, month, calendar.monthrange(year, month)[1])
			if on_rolls(emp.get("date_of_joining"), emp.get("left_date"), m_start, m_end):
				month_head[(month, sex)] += 1
		working.append(
			{
				"employee": code,
				"employee_name": emp.get("full_name") or "",
				"designation": emp.get("designation") or "",
				"date_of_joining": _as_date(emp.get("date_of_joining")),
				"months": months,
				"total": worked,
			}
		)

	manpower = []
	for month in range(1, 13):
		total_days = calendar.monthrange(year, month)[1]
		m_start = date(year, month, 1)
		m_end = date(year, month, total_days)
		closing = sum(1 for day in weekly_off_dates if m_start <= day <= m_end)
		closing += sum(1 for day in holiday_dates if m_start <= day <= m_end)
		male = month_head[(month, "male")]
		female = month_head[(month, "female")]
		male_days = month_mandays[(month, "male")]
		female_days = month_mandays[(month, "female")]
		manpower.append(
			{
				"month": month,
				"month_name": MONTHS[month - 1],
				"closing_days": closing,
				"total_days": total_days,
				"working_days": total_days - closing,
				"male": male,
				"female": female,
				"employees": male + female,
				"male_mandays": male_days,
				"female_mandays": female_days,
				"mandays": male_days + female_days,
				"ot": month_ot[month],
			}
		)

	weeks = operation_days / 7.0 if operation_days else 0
	employment = []
	for band, _label in BANDS:
		for sex, _sex_label in SEXES:
			stat = groups[(band, sex)]
			mandays = stat["mandays"]
			average = round(mandays / operation_days) if operation_days else 0
			man_hours = mandays * HOURS_PER_DAY + stat["ot"]
			per_week = round((man_hours / weeks) / average) if weeks and average else 0
			employment.append(
				{
					"band": band,
					"sex": sex,
					"mandays": mandays,
					"average_workers": average,
					"ot_hours": stat["ot"],
					"man_hours": man_hours,
					"hours_per_week": per_week,
				}
			)

	left_rows = []
	for emp in people:
		left = _as_date(emp.get("left_date"))
		if not left or left.year != year:
			continue
		left_rows.append(
			{
				"employee": emp.get("name"),
				"employee_name": emp.get("full_name") or "",
				"designation": emp.get("designation") or "",
				"date_of_joining": _as_date(emp.get("date_of_joining")),
				"left_date": left,
				"month": left.month,
			}
		)
	left_rows.sort(key=lambda row: (row["left_date"], row["employee"] or ""))

	return {
		"year": year,
		"working": working,
		"groups": groups,
		"manpower": manpower,
		"employment": employment,
		"plant": {
			"operation_days": operation_days,
			"weeks": weeks,
			"hours_per_day": HOURS_PER_DAY,
			"year_days": year_days,
			"weekly_offs": len(weekly_off_dates),
			"holidays": len(holiday_dates),
		},
		"left": left_rows,
	}


def _working_columns(year):
	columns = [
		{"label": _("Sr"), "fieldname": "sr_no", "fieldtype": "Int", "width": 50},
		{"label": _("Emp. Code"), "fieldname": "employee", "fieldtype": "Link", "options": "Company Link", "width": 120},
		{"label": _("Name of Employees"), "fieldname": "employee_name", "fieldtype": "Data", "width": 180},
		{"label": _("Designation"), "fieldname": "designation", "fieldtype": "Data", "width": 140},
		{"label": _("Date of Joining"), "fieldname": "date_of_joining", "fieldtype": "Date", "width": 120},
	]
	for index, name in enumerate(MONTHS, start=1):
		columns.append(
			{
				"label": f"{name[:3]}-{year}",
				"fieldname": f"m{index}",
				"fieldtype": "Float",
				"precision": 1,
				"width": 80,
			}
		)
	columns.append(
		{"label": _("Total"), "fieldname": "total", "fieldtype": "Float", "precision": 1, "width": 90}
	)
	return columns


def _working_rows(pack):
	rows = []
	totals = [0.0] * 12
	grand = 0.0
	for index, emp in enumerate(pack["working"], start=1):
		row = {
			"sr_no": index,
			"employee": emp["employee"],
			"employee_name": emp["employee_name"],
			"designation": emp["designation"],
			"date_of_joining": emp["date_of_joining"],
			"total": _round_day(emp["total"]),
		}
		for month, value in enumerate(emp["months"], start=1):
			shown = _round_day(value)
			row[f"m{month}"] = shown
			totals[month - 1] += value
		grand += emp["total"]
		rows.append(row)
	if rows:
		total_row = {
			"sr_no": "",
			"employee": "",
			"employee_name": _("Total"),
			"designation": "",
			"date_of_joining": None,
			"total": _round_day(grand),
			"_is_total": 1,
		}
		for month, value in enumerate(totals, start=1):
			total_row[f"m{month}"] = _round_day(value)
		rows.append(total_row)
	return rows


def _load(company, year):
	start = date(year, 1, 1)
	end = date(year, 12, 31)
	employees = frappe.db.sql(
		"""
		SELECT cl.name, cl.full_name, cl.designation, cl.date_of_joining, cl.left_date,
			cl.category, cl.weekly_off, e.gender, e.date_of_birth
		FROM `tabCompany Link` cl
		LEFT JOIN `tabEmployee` e ON e.name = cl.employee
		WHERE cl.company = %s
			AND cl.date_of_joining <= %s
			AND (cl.left_date IS NULL OR cl.left_date = '' OR cl.left_date >= %s)
		""",
		(company, end, start),
		as_dict=True,
	)
	codes = [row.name for row in employees]
	attendance = []
	daily = []
	if codes:
		attendance = frappe.db.sql(
			"""
			SELECT employee, attendance_date, status
			FROM `tabAttendance`
			WHERE company = %s AND attendance_date BETWEEN %s AND %s
			""",
			(company, start, end),
			as_dict=True,
		)
		daily = frappe.db.sql(
			"""
			SELECT employee, attendance_date, status, ot_hours
			FROM `tabAttendance Daily Rate`
			WHERE attendance_date BETWEEN %s AND %s AND employee IN %s
			""",
			(start, end, tuple(codes)),
			as_dict=True,
		)
	return assemble(employees, attendance, daily, year)


def _workbook(pack, company, year):
	wb = Workbook()
	_sheet_working(wb.active, pack, company, year)
	_sheet_leave(wb.create_sheet("Leave & Wages"), pack, company, year)
	_sheet_manpower(wb.create_sheet("Total Manpower"), pack, company, year)
	_sheet_employment(wb.create_sheet("Employment Information"), pack, company, year)
	_sheet_left(wb.create_sheet("Left Workers"), pack, company, year)
	return wb


def _sheet_working(ws, pack, company, year):
	ws.title = "Working Days"
	headers = [
		"SR. NO.",
		"EMP. CODE",
		"NAME OF EMPLOYEES",
		"DESIGNATION",
		"DATE OF JOINING",
		*[f"{name[:3]}-{year}" for name in MONTHS],
		"TOTAL",
	]
	_title(ws, 1, len(headers), f"{company} — Annual Return Form — Working Days — {year}")
	_header_row(ws, 3, headers)
	for col in range(1, len(headers) + 1):
		ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
		ws.cell(4, col).fill = HEADER_FILL
		ws.cell(4, col).border = THIN
	for row_index, emp in enumerate(pack["working"], start=5):
		values = [
			row_index - 4,
			emp["employee"],
			emp["employee_name"],
			emp["designation"],
			emp["date_of_joining"],
			*[_round_day(value) for value in emp["months"]],
			_round_day(emp["total"]),
		]
		_body_row(ws, row_index, values, date_cols={5})
	total_row = 5 + len(pack["working"])
	totals = [0.0] * 12
	grand = 0.0
	for emp in pack["working"]:
		for index, value in enumerate(emp["months"]):
			totals[index] += value
		grand += emp["total"]
	_body_row(
		ws,
		total_row,
		["", "", "Total", "", "", *[_round_day(value) for value in totals], _round_day(grand)],
		bold=True,
		fill=TOTAL_FILL,
	)
	widths = [10, 16, 28, 22, 16] + [12] * 12 + [12]
	_widths(ws, widths)
	ws.row_dimensions[3].height = 30
	ws.freeze_panes = "A5"
	ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(total_row, 4)}"
	ws.page_setup.orientation = "landscape"
	ws.page_setup.fitToPage = True
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 0
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.page_setup.paperSize = ws.PAPERSIZE_A4
	ws.oddHeader.left.text = f"{company} — Annual Return Form {year}"
	ws.oddFooter.center.text = "Working Days  |  Page &P of &N"


def _sheet_leave(ws, pack, company, year):
	headers = [
		"Workers' Category",
		"Sex",
		"Number of workers on duty during the year",
		"Number of workers eligible for annual leave",
		"No of workers whose leaves sanctioned",
		"No of workers suspended",
		"No of workers Terminated",
		"No workers left the job",
		"No of workers retired",
		"No of workers died while in service",
		"No of workers to whom payment is made instead of leave",
	]
	_title(ws, 1, len(headers), f"{company} — Leave & Wages Details — {year}")
	note = ws.cell(2, 1, "(Excluding Agriculture Team)")
	note.font = BOLD_FONT
	_header_row(ws, 3, headers)
	for index in range(1, len(headers) + 1):
		cell = ws.cell(4, index, f"({index})")
		cell.font = SUB_FONT
		cell.fill = SUB_FILL
		cell.alignment = CENTER
		cell.border = THIN
	row = 6
	for band, label in BANDS:
		start = row
		for sex, sex_label in SEXES:
			stat = pack["groups"][(band, sex)]
			_body_row(
				ws,
				row,
				[
					label if sex == "male" else "",
					sex_label,
					stat["on_duty"],
					stat["eligible"],
					stat["sanctioned"],
					0,
					0,
					stat["left"],
					0,
					0,
					0,
				],
			)
			row += 1
		ws.merge_cells(start_row=start, start_column=1, end_row=start + 1, end_column=1)
		ws.cell(start, 1).alignment = LEFT
	_widths(ws, [36, 12, 28, 28, 28, 22, 22, 22, 20, 24, 32])
	ws.row_dimensions[3].height = 48
	ws.page_setup.orientation = "landscape"
	ws.page_setup.fitToPage = True
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 1
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.page_setup.paperSize = ws.PAPERSIZE_A4
	ws.print_title_rows = "1:4"
	ws.oddFooter.center.text = "Leave & Wages  |  Page &P of &N"


def _sheet_manpower(ws, pack, company, year):
	_title(ws, 1, 12, f"{company} — Total Manpower — {year}")
	top = [
		(1, "S.N."),
		(2, "Month"),
		(3, "Closing Days"),
		(4, "Total Days"),
		(5, "Total Working Days"),
		(6, "No. of Employees"),
		(9, "Mandays"),
		(12, "OT"),
	]
	for col, label in top:
		cell = ws.cell(3, col, label)
		cell.font = HEADER_FONT
		cell.fill = HEADER_FILL
		cell.alignment = CENTER
		cell.border = THIN
	ws.merge_cells("A3:A4")
	ws.merge_cells("B3:B4")
	ws.merge_cells("C3:C4")
	ws.merge_cells("D3:D4")
	ws.merge_cells("E3:E4")
	ws.merge_cells("F3:H3")
	ws.merge_cells("I3:K3")
	ws.merge_cells("L3:L4")
	for col, label in ((6, "Male"), (7, "Female"), (8, "Total"), (9, "Male"), (10, "Female"), (11, "Total")):
		cell = ws.cell(4, col, label)
		cell.font = HEADER_FONT
		cell.fill = HEADER_FILL
		cell.alignment = CENTER
		cell.border = THIN
	for col in (1, 2, 3, 4, 5, 12):
		ws.cell(4, col).fill = HEADER_FILL
		ws.cell(4, col).border = THIN
	totals = Counter()
	for offset, month in enumerate(pack["manpower"]):
		values = [
			offset + 1,
			f"{month['month_name'][:3]}-{year}",
			month["closing_days"],
			month["total_days"],
			month["working_days"],
			month["male"],
			month["female"],
			month["employees"],
			_round_day(month["male_mandays"]),
			_round_day(month["female_mandays"]),
			_round_day(month["mandays"]),
			_round_day(month["ot"]),
		]
		_body_row(ws, 5 + offset, values)
		for key in (
			"closing_days",
			"total_days",
			"working_days",
			"male",
			"female",
			"employees",
			"male_mandays",
			"female_mandays",
			"mandays",
			"ot",
		):
			totals[key] += month[key]
	_body_row(
		ws,
		17,
		[
			"",
			"Total",
			totals["closing_days"],
			totals["total_days"],
			totals["working_days"],
			"",
			"",
			"",
			_round_day(totals["male_mandays"]),
			_round_day(totals["female_mandays"]),
			_round_day(totals["mandays"]),
			_round_day(totals["ot"]),
		],
		bold=True,
		fill=TOTAL_FILL,
	)
	_widths(ws, [8, 14, 14, 12, 18, 12, 12, 12, 12, 12, 12, 12])
	ws.row_dimensions[3].height = 22
	ws.page_setup.orientation = "landscape"
	ws.page_setup.fitToPage = True
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 1
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.page_setup.paperSize = ws.PAPERSIZE_A4
	ws.print_title_rows = "1:4"
	ws.oddFooter.center.text = "Total Manpower  |  Page &P of &N"


def _sheet_employment(ws, pack, company, year):
	plant = pack["plant"]
	_title(ws, 1, 7, f"{company} — Employment Information — {year}")
	facts = [
		(3, "1", "No of days plant was in operation", plant["operation_days"], "Total days of year", plant["year_days"]),
		(4, "2", "No of weeks plant was in operation", _round_day(plant["weeks"]), "Total W. Off", plant["weekly_offs"]),
		(5, "3", "Total hours per day workers working (excluding overtime)", plant["hours_per_day"], "Total Holidays", plant["holidays"]),
		(6, "", "", "", "Total days of operation", plant["operation_days"]),
	]
	for row, serial, label, value, side_label, side_value in facts:
		ws.cell(row, 1, serial).font = BOLD_FONT
		ws.cell(row, 2, label).font = BODY_FONT
		ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
		num = ws.cell(row, 5, value)
		num.font = BOLD_FONT
		num.alignment = CENTER
		ws.cell(row, 6, side_label).font = BODY_FONT
		side = ws.cell(row, 7, side_value)
		side.font = BOLD_FONT
		side.alignment = CENTER
	ws.cell(8, 1, "Number of labors and their details").font = TITLE_FONT
	headers = [
		"Category of workers",
		"Sex",
		"No. of Mandays during the year",
		"Average number of Workers employed per day",
		"Number of man-hours worked on overtime in a year",
		"Number of manhours worked including overtime in a year",
		"Average number of hours worked per week",
	]
	_header_row(ws, 10, headers)
	for index in range(1, 8):
		cell = ws.cell(11, index, f"({index})")
		cell.font = SUB_FONT
		cell.fill = SUB_FILL
		cell.alignment = CENTER
		cell.border = THIN
	by_key = {(row["band"], row["sex"]): row for row in pack["employment"]}
	row_no = 12
	for band, label in BANDS:
		start = row_no
		for sex, sex_label in SEXES:
			stat = by_key[(band, sex)]
			_body_row(
				ws,
				row_no,
				[
					label if sex == "male" else "",
					sex_label,
					_round_day(stat["mandays"]),
					stat["average_workers"],
					_round_day(stat["ot_hours"]),
					_round_day(stat["man_hours"]),
					stat["hours_per_week"],
				],
			)
			row_no += 1
		ws.merge_cells(start_row=start, start_column=1, end_row=start + 1, end_column=1)
	_widths(ws, [36, 12, 24, 28, 32, 36, 28])
	ws.row_dimensions[10].height = 36
	ws.page_setup.orientation = "landscape"
	ws.page_setup.fitToPage = True
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 1
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.page_setup.paperSize = ws.PAPERSIZE_A4
	ws.print_title_rows = "1:11"
	ws.oddFooter.center.text = "Employment Information  |  Page &P of &N"


def _sheet_left(ws, pack, company, year):
	headers = [
		"SR.NO.",
		"EMP. CODE",
		"NAME OF WORKER",
		"DESIGNATION",
		"DATE OF JOINING",
		"Exit date",
		*[f"{name[:3]}-{year}" for name in MONTHS],
		"Total",
	]
	_title(ws, 1, len(headers), f"{company} — Left Workers — {year}")
	_header_row(ws, 3, headers)
	for col in range(1, len(headers) + 1):
		ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
		ws.cell(4, col).fill = HEADER_FILL
		ws.cell(4, col).border = THIN
	if not pack["left"]:
		ws.cell(5, 1, "No workers left in this year").font = BODY_FONT
	for index, emp in enumerate(pack["left"], start=1):
		months = [1 if month == emp["month"] else 0 for month in range(1, 13)]
		_body_row(
			ws,
			4 + index,
			[
				index,
				emp["employee"],
				emp["employee_name"],
				emp["designation"],
				emp["date_of_joining"],
				emp["left_date"],
				*months,
				1,
			],
			date_cols={5, 6},
		)
	_widths(ws, [10, 16, 28, 22, 16, 16] + [12] * 12 + [10])
	ws.freeze_panes = "A5"
	ws.page_setup.orientation = "landscape"
	ws.page_setup.fitToPage = True
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 0
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.page_setup.paperSize = ws.PAPERSIZE_A4
	ws.print_title_rows = "1:4"
	last = 4 + max(len(pack["left"]), 1)
	ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last}"
	ws.oddFooter.center.text = "Left Workers  |  Page &P of &N"


def _title(ws, row, columns, text):
	ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
	cell = ws.cell(row, 1, text)
	cell.font = TITLE_FONT
	cell.alignment = Alignment(horizontal="left", vertical="center")
	ws.row_dimensions[row].height = 22


def _header_row(ws, row, headers):
	for col, label in enumerate(headers, start=1):
		cell = ws.cell(row, col, label)
		cell.font = HEADER_FONT
		cell.fill = HEADER_FILL
		cell.alignment = CENTER
		cell.border = THIN


def _body_row(ws, row, values, date_cols=(), bold=False, fill=None):
	for col, value in enumerate(values, start=1):
		cell = ws.cell(row, col, value if value != "" else None)
		cell.font = BOLD_FONT if bold else BODY_FONT
		cell.border = THIN
		cell.alignment = LEFT if col in {2, 3, 4} else CENTER
		if fill is not None:
			cell.fill = fill
		if col in date_cols and isinstance(value, date):
			cell.number_format = "DD-MM-YYYY"
		elif isinstance(value, float):
			cell.number_format = "0.0"


def _widths(ws, widths):
	for index, width in enumerate(widths, start=1):
		ws.column_dimensions[get_column_letter(index)].width = width


def _save_excel(wb, prefix):
	buf = io.BytesIO()
	wb.save(buf)
	stamp = frappe.utils.now_datetime().strftime("%Y%m%d_%H%M%S")
	fn = f"{prefix}_{stamp}.xlsx"
	with open(frappe.utils.get_files_path(fn, is_private=0), "wb") as handle:
		handle.write(buf.getvalue())
	doc = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": fn,
			"is_private": 0,
			"file_url": f"/files/{fn}",
		}
	)
	doc.insert(ignore_permissions=True)
	frappe.db.commit()
	return doc.file_url


def _normalize_filters(filters):
	if isinstance(filters, str):
		filters = json.loads(filters or "{}")
	return frappe._dict(filters or {})


def _year(filters):
	year = cint(filters.get("year"))
	if year < 1900:
		return 0
	return year


def _as_date(value):
	if not value:
		return None
	if isinstance(value, datetime):
		return value.date()
	if isinstance(value, date):
		return value
	return getdate(value)


def _round_day(value):
	number = round(flt(value), 1)
	if abs(number - round(number)) < 1e-9:
		return int(round(number))
	return number


def _common_weekly_off(employees):
	counts = Counter((emp.get("weekly_off") or "").strip() for emp in employees)
	counts.pop("", None)
	if not counts:
		return WEEKDAYS["Sunday"]
	name = counts.most_common(1)[0][0]
	return WEEKDAYS.get(name, WEEKDAYS["Sunday"])


def _next_day(day):
	return date.fromordinal(day.toordinal() + 1)
